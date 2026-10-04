"""Regenerate the PDF and DOCX sample documents (the .md and .txt samples are edited directly).

Run inside the backend container, which has fpdf2 and python-docx:
    docker compose run --rm --no-deps -v "./samples:/samples" backend python /samples/build_samples.py
"""

from pathlib import Path

import docx
from fpdf import FPDF

OUT = Path(__file__).parent / "docs"

HANDBOOK = [
    (
        "Installation Handbook",
        [
            "This handbook describes how Northwind Solar crews install rooftop systems. Every job "
            "starts with a structural survey of the roof, carried out by a certified technician "
            "from the Bergen workshop or the Rotterdam office.",
            "The survey checks rafter spacing, roof age and load capacity. A roof older than 25 "
            "years must be inspected by a structural engineer before panels are mounted.",
        ],
    ),
    (
        "Shading analysis",
        [
            "A shading analysis estimates the annual energy yield from drone images and local "
            "weather data. If any panel position is shaded for more than two hours a day, the "
            "crew fits Tidewater Micro microinverters instead of a single Tidewater string "
            "inverter.",
        ],
    ),
    (
        "Mounting",
        [
            "NW-400 panels are mounted on anodised aluminium rails bolted into the roof rafters. "
            "Every bolt is sealed with butyl tape to keep water out. On coastal jobs all "
            "fasteners are stainless steel, because salt spray is the main cause of corrosion.",
            "Rails are spaced 1.2 metres apart. The crew torques each bolt to 25 newton-metres "
            "and records the value in the job sheet.",
        ],
    ),
    (
        "Electrical work",
        [
            "Inverters from Kestrel Power Systems are installed in a shaded, ventilated location. "
            "The Fjord LFP-10 battery, when ordered, goes in a ventilated cabinet away from direct "
            "sunlight, because high temperatures shorten battery life.",
            "Commissioning ends with a two-hour monitored test run. The inverter logs are uploaded "
            "to the customer portal before the crew leaves the site.",
        ],
    ),
    (
        "Annual maintenance",
        [
            "Technicians visit once a year. They clean salt deposits from the panel glass, check "
            "that the rail bolts are still at 25 newton-metres, and download inverter logs to find "
            "underperforming strings.",
        ],
    ),
]

WARRANTY = [
    ("Warranty and Service Terms", None),
    (
        "Panel warranty",
        "NW-400 panels are covered for twenty-five years against output loss below eighty percent "
        "of the rated power. Physical damage from storms is covered by the customer's building "
        "insurance, not by the panel warranty.",
    ),
    (
        "Inverter warranty",
        "Tidewater inverters and Tidewater Micro microinverters carry a ten-year warranty from "
        "Kestrel Power Systems, which Northwind Solar passes on to the customer. Northwind handles "
        "the claim and the replacement on the customer's behalf.",
    ),
    (
        "Battery warranty",
        "The Fjord LFP-10 battery from Aster Cells is guaranteed to keep seventy percent of its "
        "capacity for ten years or 6,000 cycles, whichever comes first.",
    ),
    (
        "Workmanship guarantee",
        "Workmanship, including roof leaks caused by the installation, is guaranteed for five "
        "years. Leased systems are covered for the whole fifteen-year lease term.",
    ),
    (
        "Making a claim",
        "Claims are filed through the customer portal with photos of the fault. Priya Raman's "
        "customer service team confirms receipt within two working days and books a technician "
        "visit within ten working days.",
    ),
]


def build_pdf(path: Path) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    for title, paragraphs in HANDBOOK:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.multi_cell(0, 9, title)
        pdf.ln(2)
        pdf.set_font("Helvetica", size=11)
        for p in paragraphs:
            pdf.multi_cell(0, 6, p)
            pdf.ln(3)
    pdf.output(str(path))


def build_docx(path: Path) -> None:
    document = docx.Document()
    for heading, body in WARRANTY:
        document.add_heading(heading, level=1 if body is None else 2)
        if body:
            document.add_paragraph(body)
    document.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_pdf(OUT / "installation_handbook.pdf")
    build_docx(OUT / "warranty_terms.docx")
    print("wrote", sorted(p.name for p in OUT.iterdir()))
