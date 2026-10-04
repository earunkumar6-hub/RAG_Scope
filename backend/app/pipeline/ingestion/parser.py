"""S2 parse & clean: per-page text extraction, whitespace normalisation, header/footer removal."""

import io
import re
from collections import Counter
from dataclasses import dataclass
from html.parser import HTMLParser

PAGE_SEPARATOR = "\n\n"
HEADER_FOOTER_LINES = 2  # candidate lines inspected at the top and bottom of each page
MIN_PAGES_FOR_HEADER_FOOTER = 3


@dataclass(frozen=True)
class PageSpan:
    page: int  # 1-based
    start: int  # char offset in ParsedDocument.text
    end: int


@dataclass
class ParsedDocument:
    text: str
    pages: list[PageSpan]
    page_count: int  # physical pages in the source (including empty ones)
    raw_chars: int
    raw_preview: str
    removed_lines: list[str]

    def page_at(self, offset: int) -> int:
        """Page number containing char ``offset`` (offsets in separators map to the next page)."""
        for span in self.pages:
            if offset < span.end:
                return span.page
        return self.pages[-1].page if self.pages else 1


class _TextExtractor(HTMLParser):
    BLOCK_TAGS = {"p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "tr", "br", "hr"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag in self.BLOCK_TAGS:
            self.parts.append("\n\n" if tag != "br" else "\n")

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


def extract_pages(extension: str, content: bytes) -> list[str]:
    """Raw text per page. DOCX/MD/TXT have no physical pages: form feeds split pages, else one."""
    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return [page.extract_text() or "" for page in reader.pages]
    if extension == ".docx":
        import docx

        document = docx.Document(io.BytesIO(content))
        blocks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text for cell in row.cells))
        return ["\n\n".join(blocks)]
    text = content.decode("utf-8-sig")
    if extension == ".md":
        import markdown

        extractor = _TextExtractor()
        extractor.feed(markdown.markdown(text, extensions=["tables", "fenced_code"]))
        text = "".join(extractor.parts)
    return text.split("\f") if "\f" in text else [text]


def count_pages(extension: str, content: bytes) -> int:
    """Physical page count (cheap; also proves the file opens)."""
    if extension == ".pdf":
        from pypdf import PdfReader

        return len(PdfReader(io.BytesIO(content)).pages)
    if extension == ".docx":
        import docx

        docx.Document(io.BytesIO(content))
        return 1
    return content.decode("utf-8-sig").count("\f") + 1


def normalise(text: str) -> str:
    """Normalise whitespace while keeping paragraph breaks (used by the chunker as boundaries)."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate words broken across lines
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _line_key(line: str) -> str:
    """Comparison key for repeated lines: page numbers / digits do not matter."""
    return re.sub(r"\d+", "#", line.strip().lower())


def remove_repeated_headers_footers(pages: list[str]) -> tuple[list[str], list[str]]:
    """Drop top/bottom lines whose (digit-insensitive) text repeats on more than 50% of pages."""
    if len(pages) < MIN_PAGES_FOR_HEADER_FOOTER:
        return pages, []
    page_lines = [[ln for ln in p.split("\n") if ln.strip()] for p in pages]
    counts: Counter[str] = Counter()
    for lines in page_lines:
        edge = lines[:HEADER_FOOTER_LINES] + lines[-HEADER_FOOTER_LINES:]
        counts.update({_line_key(ln) for ln in edge})
    repeated = {key for key, n in counts.items() if n > len(pages) / 2 and key}
    if not repeated:
        return pages, []
    removed: dict[str, str] = {}  # key -> first example ("Page 1" stands for "Page #")
    cleaned: list[str] = []
    for lines in page_lines:
        n = len(lines)
        keep = []
        for i, ln in enumerate(lines):
            at_edge = i < HEADER_FOOTER_LINES or i >= n - HEADER_FOOTER_LINES
            if at_edge and _line_key(ln) in repeated:
                removed.setdefault(_line_key(ln), ln.strip())
                continue
            keep.append(ln)
        cleaned.append("\n".join(keep))
    return cleaned, list(removed.values())[:20]


def parse_document(extension: str, content: bytes, preview_chars: int = 1500) -> ParsedDocument:
    """Extract, clean and concatenate pages; returns clean text plus a page map."""
    raw_pages = extract_pages(extension, content)
    raw_joined = PAGE_SEPARATOR.join(raw_pages)
    pages, removed = remove_repeated_headers_footers([p.replace("\r\n", "\n") for p in raw_pages])
    spans: list[PageSpan] = []
    parts: list[str] = []
    offset = 0
    for number, page_text in enumerate(pages, start=1):
        clean = normalise(page_text)
        if not clean:
            continue
        if parts:
            parts.append(PAGE_SEPARATOR)
            offset += len(PAGE_SEPARATOR)
        spans.append(PageSpan(page=number, start=offset, end=offset + len(clean)))
        parts.append(clean)
        offset += len(clean)
    return ParsedDocument(
        text="".join(parts),
        pages=spans,
        page_count=len(raw_pages),
        raw_chars=len(raw_joined),
        raw_preview=raw_joined[:preview_chars],
        removed_lines=removed,
    )
