"""Test doubles and fixture-file builders (no model downloads, no network)."""

import hashlib
import io
import json
import random
import re

import numpy as np

from app.llm.base import LLMError, LLMProvider, LLMResult

TOPICS = {
    "graphs": [
        "vertex",
        "edge",
        "traversal",
        "neighbour",
        "adjacency",
        "path",
        "hop",
        "entity",
        "relation",
        "knowledge",
    ],
    "vectors": [
        "embedding",
        "cosine",
        "similarity",
        "dimension",
        "index",
        "nearest",
        "neighbour",
        "vector",
        "space",
    ],
    "cooking": [
        "flour",
        "butter",
        "oven",
        "recipe",
        "bake",
        "sugar",
        "dough",
        "knead",
        "yeast",
        "crust",
    ],
    "finance": [
        "invoice",
        "ledger",
        "revenue",
        "audit",
        "balance",
        "credit",
        "debit",
        "account",
        "payment",
        "tax",
    ],
}


def make_text(seed: int = 7, paragraphs: int = 24, topic_cycle: list[str] | None = None) -> str:
    """Deterministic multi-topic prose with paragraph and sentence boundaries."""
    rng = random.Random(seed)
    topics = topic_cycle or list(TOPICS)
    paras = []
    for p in range(paragraphs):
        vocab = TOPICS[topics[p % len(topics)]]
        sentences = []
        for _ in range(rng.randint(3, 6)):
            words = [rng.choice(vocab) for _ in range(rng.randint(8, 16))]
            sentences.append(" ".join(words).capitalize() + ".")
        paras.append(" ".join(sentences))
    return "\n\n".join(paras)


class FakeEmbedder:
    """Deterministic hashed bag-of-words embedder implementing the ``Embedder`` protocol."""

    model_name = "fake-hash-64"

    def __init__(self, dim: int = 64, fail: bool = False):
        self._dim = dim
        self.fail = fail
        self.loaded = True

    @property
    def dimension(self) -> int:
        return self._dim

    def load(self) -> None:
        return None

    def _vec(self, text: str) -> np.ndarray:
        v = np.zeros(self._dim, dtype=np.float32)
        for word in re.findall(r"\w+", text.lower()):
            v[int(hashlib.md5(word.encode()).hexdigest(), 16) % self._dim] += 1.0
        n = np.linalg.norm(v)
        return v / n if n else v

    def embed_documents(self, texts, on_batch=None):  # noqa: ANN001, ANN201
        if self.fail:
            raise RuntimeError("embedding backend exploded")
        out = np.vstack([self._vec(t) for t in texts]) if texts else np.zeros((0, self._dim))
        if on_batch:
            on_batch(len(texts), len(texts))
        return out

    def embed_query(self, text: str) -> np.ndarray:
        return self._vec(text)

    def count_truncated(self, texts: list[str]) -> tuple[int, int]:
        return 0, 10_000


def make_pdf(pages: list[str], header: str | None = None, footer: bool = True) -> bytes:
    """Multi-page PDF with an optional repeated header and 'Page N' footer."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=False)
    for number, body in enumerate(pages, start=1):
        pdf.add_page()
        pdf.set_font("Helvetica", size=10)
        if header:
            pdf.set_xy(10, 8)
            pdf.cell(0, 6, header)
        pdf.set_xy(10, 20)
        pdf.multi_cell(0, 5, body)
        if footer:
            pdf.set_xy(10, 285)
            pdf.cell(0, 6, f"Page {number}")
    return bytes(pdf.output())


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def parse_sse(body: str) -> list[dict]:
    """Parse an SSE body into [{id, event, data}] (data JSON-decoded)."""
    events = []
    for block in re.split(r"\r?\n\r?\n", body):
        fields: dict[str, str] = {}
        for line in block.splitlines():
            if not line or line.startswith(":"):
                continue
            key, _, value = line.partition(":")
            fields[key] = value[1:] if value.startswith(" ") else value
        if "data" in fields:
            events.append(
                {
                    "id": fields.get("id"),
                    "event": fields.get("event"),
                    "data": json.loads(fields["data"]),
                }
            )
    return events


class FakeReranker:
    """Scores by query-word overlap (deterministic, no model)."""

    model_name = "fake-overlap"
    loaded = True

    def score(self, query: str, texts: list[str]) -> list[float]:
        q = set(re.findall(r"\w+", query.lower()))
        return [len(q & set(re.findall(r"\w+", t.lower()))) / (len(q) or 1) for t in texts]


VOCAB = {w for words in TOPICS.values() for w in words}


def _vocab_words(text: str, limit: int) -> list[str]:
    seen: list[str] = []
    for w in re.findall(r"[a-z]+", text.lower()):
        if w in VOCAB and w not in seen:
            seen.append(w)
    return seen[:limit]


class FakeLLM(LLMProvider):
    """Records calls; streams an answer citing the first context chunk (plus ``extra``).

    ``complete`` answers S8 extraction (entities = known topic words in the passage, chained by
    "related to") and Q5 query-entity extraction; anything else (e.g. cluster labels) fails like
    an unavailable provider, so ingestion falls back to keyword labels.
    """

    name = "fake"
    model = "fake-llm"
    fast_model = "fake-llm"

    def __init__(self, extra: str = ""):
        self.extra = extra
        self.calls: list[list[dict[str, str]]] = []
        self.complete_calls: list[list[dict[str, str]]] = []

    def complete(self, messages, **kwargs):  # noqa: ANN001, ANN201
        system, user = messages[0]["content"], messages[-1]["content"]
        self.complete_calls.append(messages)
        if "knowledge graph from one passage" in system:
            words = _vocab_words(user, 6)
            out = {
                "entities": [
                    {"name": w.title(), "type": "CONCEPT", "description": f"{w} concept"}
                    for w in words
                ],
                "relations": [
                    {"source": a.title(), "relation": "related to", "target": b.title()}
                    for a, b in zip(words, words[1:], strict=False)
                ],
            }
        elif "entities mentioned in the question" in system:
            out = {"entities": [w.title() for w in _vocab_words(user, 5)]}
        else:
            raise LLMError("fake LLM: unsupported completion")
        return LLMResult(json.dumps(out), self.model, 50, 20)

    def stream(self, messages, on_token, **kwargs):  # noqa: ANN001, ANN201
        self.calls.append(messages)
        first_id = re.search(r"^\[([^\]]+)\]", messages[1]["content"], re.M).group(1)
        parts = ["The answer ", "is here ", f"[{first_id}].", self.extra]
        for p in parts:
            if p:
                on_token(p)
        return LLMResult("".join(parts), self.model, 100, 10)
